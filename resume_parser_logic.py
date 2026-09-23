import json
import os
from pathlib import Path
from dotenv import load_dotenv
from groq import Groq
from pydantic import BaseModel, Field


load_dotenv()
my_api_key=os.getenv("GROQ_API_KEY")
if not my_api_key:
    raise ValueError("API KEY KAHA HAI BHAI")

client =Groq(api_key=my_api_key)
model="openai/gpt-oss-120b"

class Education(BaseModel):
    college:str
    degree:str
    graduation_year:int
    cgpa:float | None=None
    
class Experience(BaseModel):
    company: str | None=None
    role: str | None=None
    duration: str | None=None
    description:str | None=None
    skills_used:list[str]=Field(default_factory=list)
    
class Project(BaseModel):
    name:str
    type:str
    description:str
    technologies:list[str]=Field(default_factory=list)
    responsibilities:list[str]= Field(default_factory=list)

class Internship(BaseModel):
    company:str
    role:str
    duration:str
    description:str
    skills_used:list[str]=Field(default_factory=list)

class HRProfile(BaseModel):
    reason_for_leaving: str | None = None
    strengths: list[str] = Field(default_factory=list)
    weaknesses: list[str] = Field(default_factory=list)     
        
class Resume(BaseModel):
    name:str | None=None
    email: str | None=None
    phone: str |None=None    
    total_experience: float | None=None
    skills:list[str]=Field(default_factory=list)
    experience:list[Experience]=Field(default_factory=list)
    education:list[Education]=Field(default_factory=list)
    projects:list[Project]=Field(default_factory=list)
    certifications:list[str]=Field(default_factory=list)
    professional_summary:str | None=None
    github:str | None=None
    linkedin:str| None=None
    internships:list[Internship]=Field(default_factory=list)
    hr_profile: HRProfile | None = None

class JDMatchingResult(BaseModel):
    match_percentage:float
    matched_skills: list[str] = Field(default_factory=list)
    missing_skills: list[str] = Field(default_factory=list)
    relevant_experience:str | None=None
    relevant_projects:str |None=None
    explanation: str

   

    
resume_schema=Resume.model_json_schema() 
match_result=JDMatchingResult.model_json_schema()   

def match_jd(jd_text,candidate_profile):
    system_prompt=f"""
    You are an expert HR assistant.

    TASK
    Compare the job description with the candidate profile
    and determine how well the candidate matches the requirements.

    CANDIDATE PROFILE
    {candidate_profile}

    JOB DESCRIPTION
    {jd_text}

    SOURCE OF TRUTH
    The candidate profile is the single source of truth
    about the candidate.

    IMPORTANT RULES:
    1. Do not invent information.
    2. Do not infer skills that are not explicitly supported
       by the candidate profile.
    3. Do not treat related technologies as equivalent skills.
    4. Consider relevant skills, experience, projects,
       technologies, and other explicit requirements.
    5. Generate a match percentage based on the requirements
       explicitly supported by the candidate profile.
    6. Clearly identify skills/requirements that are not
       supported by the candidate profile.
    7. Return JSON matching the JDMatchingResult schema: {match_result}.
    """
    
    user_prompt = """
    Compare the provided job description with the provided candidate profile
    and generate the JD matching result.
    """
    system_message={
            "role":"system",
            "content":system_prompt
        }          
    user_message={
            "role":"user",
            "content":user_prompt
        }  
    response_format={
                "type":"json_object"
            }
    messages=[system_message,user_message]
    response=client.chat.completions.create(model=model,messages=messages,response_format=response_format)
    raw_output=response.choices[0].message.content
    data=json.loads(raw_output)
    result=JDMatchingResult(**data)
    return result.model_dump()
    
    
def parse_resume(resume_text):
    system_prompt = f"""
        You are an expert resume parser.
    
        Extract information from the resume based on its meaning,
        not only based on exact section headings.
    
        Different resumes may use different headings.
    
        For example:
        - Experience
        - Professional Experience
        - Work History
        - Employment
        - Internships
    
        These may all contain relevant experience.
    
        Skills may also appear in the skills section, work experience,
        internships or projects.
    
        Return ONLY valid JSON matching this schema:
    
        {resume_schema} 
        Important rules:
        
            1. Do not invent information.
            2. If a value is not available, return null.
            3. If a list has no information, return an empty list.
            4. Put full-time/professional employment in experience and internships in internships.
            5. Extract skills mentioned across the entire resume."""
    user_prompt = f"""
        Parse the following resume:
    
        {resume_text}
        """
    system_message={
        "role":"system",
        "content":system_prompt
    }          
    user_message={
        "role":"user",
        "content":user_prompt
    }  
    response_format={
            "type":"json_object"
        }
    messages=[system_message,user_message]
    response=client.chat.completions.create(model=model,messages=messages,response_format=response_format)
    raw_output=response.choices[0].message.content
    data=json.loads(raw_output)
    resume=Resume(**data)
    return resume.model_dump()
         
from pypdf import PdfReader
from docx import Document

def read_pdf(file_path):
    reader=PdfReader(file_path)
    text=""
    for page in reader.pages:
        page_text=page.extract_text()
        if page_text:
            text+=page_text+"\n"
    return text

def read_docx(file_path):
    document=Document(file_path)
    text=""
    for paragraph in document.paragraphs:
        if paragraph.text.strip():
            text+=paragraph.text+"\n"
    
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    text+=cell.text+"\n"
    return text

def read_resume(file_path):
    if file_path.suffix.lower()==".pdf":
        return read_pdf(file_path)
    elif file_path.suffix.lower()==".docx":
        return read_docx(file_path)
    else:
        return None

def process_resume(file_path):
    if file_path.suffix.lower() not in [".pdf", ".docx"]:
        raise ValueError("Resume must be .pdf or .docx format")

    print("\nProcessing:", file_path.name)

    try:
        resume_text = read_resume(file_path)

        if not resume_text:
            print("Could not extract resume text. Skipping...")
            return
        
        hr_profile = None

        if Path("candidate_profile.json").exists():
            with open("candidate_profile.json", "r", encoding="utf-8") as f:
                existing_profile = json.load(f)

            hr_profile = existing_profile.get("hr_profile")

        parsed_resume = parse_resume(resume_text)
        if hr_profile:
            parsed_resume["hr_profile"] = hr_profile

        with open("candidate_profile.json", "w", encoding="utf-8") as f:
            json.dump(parsed_resume, f, indent=4)
        print("Candidate profile updated successfully.")
    

    except Exception as e:
        print("ERROR processing:", file_path.name)
        print(e)


if __name__ == "__main__":
    resume_folder = Path("resumes")
    file_path = resume_folder / "Tanay_Tijare_Resume.pdf"

    process_resume(file_path)